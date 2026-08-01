#!/usr/bin/env python3
"""Check one converted file from the smoke test.

Usage: verify-conversion.py <ext> <path-to-converted-xlsx>

Exists as a file rather than a heredoc inside the workflow because a heredoc
nested in a YAML block scalar depends on the dedent landing exactly right, and
gets silently wrong the moment someone re-indents the step.

Checks three things, because an HTTP 200 proves none of them:
  1. the response really is an xlsx,
  2. it holds the rows the fixture went in with,
  3. no formula survived the conversion.

The .xlsx fixture ships with three live formulas specifically so (3) has
something to catch.
"""
import re
import sys
import zipfile

from openpyxl import load_workbook

EXPECTED = [("Alpha", 3, 9.99), ("Beta", 7, 21.5), ("Gamma", 12, 4.25)]


def main() -> int:
    if len(sys.argv) != 3:
        print(f"usage: {sys.argv[0]} <ext> <path>", file=sys.stderr)
        return 2

    ext, path = sys.argv[1], sys.argv[2]
    problems = []

    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, FileNotFoundError) as exc:
        print(f"  {ext}: not a readable xlsx -- {exc}", file=sys.stderr)
        return 1

    if "xl/workbook.xml" not in archive.namelist():
        problems.append("missing xl/workbook.xml")

    formulas = [
        match
        for name in archive.namelist()
        if name.startswith("xl/worksheets/") and name.endswith(".xml")
        for match in re.findall(r"<f[ >]", archive.read(name).decode(errors="ignore"))
    ]
    if formulas:
        problems.append(f"{len(formulas)} formula(s) survived conversion")

    sheet = load_workbook(path).active
    rows = [
        tuple(row)
        for row in sheet.iter_rows(min_row=2, max_row=4, max_col=3, values_only=True)
    ]
    if rows != EXPECTED:
        problems.append(f"rows {rows} != {EXPECTED}")

    if problems:
        for problem in problems:
            print(f"  {ext}: {problem}", file=sys.stderr)
        return 1

    print(f"  {ext}: valid xlsx, rows match, no formulas left")
    return 0


if __name__ == "__main__":
    sys.exit(main())
